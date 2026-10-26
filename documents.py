import html
def _esc(s): return html.escape("" if s is None else str(s))  # model JSON may return numbers
from questions import CARE_QUESTIONS
from forms import is_full, tick_list
def pack_html(brand, personal, draft, tier, filled_by="", prefill_off=False):
    e=_esc
    rows="".join(f"<h3>Question {q}: {e(t)}</h3><p>{e(draft['answers'].get(str(q),'')) or '<em>Not discussed. Leave blank or answer No if it does not apply.</em>'}</p>" for q,t,_,_ in CARE_QUESTIONS)
    answered=tick_list("aa",draft.get("answers",{})); model=[str(x) for x in draft.get("tick_yes",[]) if str(x) in answered]
    ticks=", ".join(model or answered)  # model's judgement limited to answered questions; else every answered question
    li=lambda xs:"".join(f"<li>{e(x)}</li>" for x in xs)
    script=""
    if is_full(tier):
        script="""<h2>Online claim: step by step</h2><ol><li>Go to gov.uk/attendance-allowance/how-to-claim and choose to apply online (if available to you) or print the filled-in AA1 (attached).</li>
<li>"""+("The filled-in AA1 already has your approved answers and the tick boxes you confirmed. Anything too long for the boxes is on extra sheets at the end of the PDF. Check every page." if not prefill_off else "GOV.UK has just published a new version of the AA1, so we have not filled it in this time. Download the latest form and copy each answer below into it.")+"""</li>
<li>Applying online instead? Keep this document open next to you, tick the box that matches each care question, and paste our prepared answer into the "tell us more" box. Questions where you described help (tick Yes): """+e(ticks)+"""</li><li>Add the NI number, bank details, signature and date yourself. We never fill these in.</li><li>Attach or post copies (never originals) of the evidence on your checklist.</li><li>Keep a copy and note the date you sent it. DWP usually writes within a few weeks.</li></ol>"""
    return f"""<!doctype html><html lang="en-GB"><meta charset="utf-8"><title>Your Attendance Allowance answers</title>
<style>body{{font:18px/1.6 system-ui;max-width:800px;margin:2rem auto;padding:0 1rem;color:#14211f}}h3{{margin-bottom:.2rem}}.box{{background:#fffaf0;border-left:5px solid #f2c14e;padding:1rem}}</style>
<h1>Your Attendance Allowance (AA1) answers</h1><p>Prepared by {e(brand)} for {e(personal.get('title',''))} {e(personal.get('other_names',''))} {e(personal.get('surname',''))}. {e(filled_by)}</p>
<div class="box"><strong>Please check every answer.</strong> It must be true for the person claiming. Change anything that isn't quite right. You send the form yourself. We can't guarantee an award.</div>
<h2>Things to check</h2><ul>{li(draft.get('notes_for_user',[]))}</ul>
<h2>Care questions</h2>{rows}<p><strong>Tick "Yes" on questions:</strong> {e(ticks)}</p>
<h2>Evidence checklist</h2><ul>{li(draft.get('checklist',[]))}</ul>
<h2>Draft letter to the GP</h2><pre style="white-space:pre-wrap;font:inherit">{e(draft.get('gp_letter',''))}</pre>
<h2>Draft statement from someone who knows them (AA1 page 28)</h2><p>{e(draft.get('carer_statement',''))}</p>{script}
<p style="font-size:.85em;color:#555">Independent service, not affiliated with DWP.</p></html>"""

from forms import FORMS, PIP_ACTIVITIES
CSS="<style>body{font:18px/1.6 system-ui;max-width:800px;margin:2rem auto;padding:0 1rem;color:#14211f}h3{margin-bottom:.2rem}.box{background:#fffaf0;border-left:5px solid #f2c14e;padding:1rem}.tip{color:#3d4f4c;font-size:.9em}</style>"
NEWVER=("<div class=box><strong>GOV.UK has just published a new version of this form</strong>, so we have not filled it in for you this time. "
        "Download the latest form from GOV.UK and copy each answer below into it.</div>")
def pack_form(brand, personal, draft, tier, form, prefill_off=False):
    if form=="aa": return pack_html(brand,personal,draft,tier,prefill_off=prefill_off)
    e=_esc; F=FORMS[form]; a=draft.get("answers",{})
    li=lambda xs:"".join(f"<li>{e(x)}</li>" for x in xs)
    nd='<em>Not discussed. If this applies, add it in your own words; if not, tick "No".</em>'
    if form=="dla":
        rows="".join(f"<h3>{e(t)}</h3><p>{e(a.get(i,'')) or nd}</p>" for i,t,_ in F["topics"])
        delivery=(NEWVER if is_full(tier) and prefill_off else "<p>Your filled-in DLA1A is attached. The answer boxes and the tick boxes you confirmed are filled. Anything too long for the boxes is on extra sheets at the end of the PDF. Check every page, then add what we never fill in (below).</p>" if is_full(tier)
          else "<p>Copy each answer into the matching question on the DLA1A form (download it from gov.uk/disability-living-allowance-children/how-to-claim).</p>")
        extra=f"<p><strong>Tick 'Yes' on questions:</strong> {e(', '.join(tick_list(form,a)) or ', '.join(map(str,draft.get('tick_yes',[]))))}</p>"
        nights="".join(f"<li>{e(k.replace('night_','').title())}: {e(v[0])} nights a week, {e(v[1])} minutes, {e(v[2])} times a night</li>" for k,v in draft.get("extra",{}).items() if k.startswith("night_"))
        if nights: extra+=f"<h3>Q70 night-time numbers</h3><ul>{nights}</ul>"
        send="<h2>Sending it</h2><ol><li>Add the child's NI number (if they have one), your NI number, bank details, signature and date yourself. We never fill these in.</li><li>Post it to Freepost DWP DLA Child (no stamp or postcode needed), with copies (never originals) of your evidence.</li><li>The claim starts when DWP receives the form, or from the date you called the helpline if you return the form within 6 weeks. Keep a copy.</li></ol>"
        title="Your Child DLA (DLA1A) answers"
    else:
        rows=f"<h3>{e(F['topics'][0][1])}</h3><p>{e(a.get('conditions','')) or nd}</p>"
        for n,name,q,_ in PIP_ACTIVITIES:
            where=f"PIP2 {q}" if form=="pip" else f"AR1 Activity {n} – 'Please tell us why and when' box"
            rows+=f"<h3>Activity {n}: {e(name)} <small>({where})</small></h3><p>{e(a.get(f'a{n}','')) or 'No difficulty described. Say so on the form if that is true.'}</p>"
        rows+=f"<h3>Additional information ({'PIP2 Q15' if form=='pip' else 'AR1 final question'})</h3><p>{e(a.get('other',''))}</p>"
        extra="<p class=tip>For each activity, describe what happens on most days (more than half), and whether you can do it safely, to a good standard, as often as you need to and in a reasonable time. Include aids and any help or prompting. We do not estimate points or outcomes.</p>"
        if form=="pipr" and draft.get("changes"): extra=f"<h3>What has changed since your last award</h3><p>{e(draft['changes'])}</p>"+extra
        delivery="<p>The PIP form is sent to you by DWP (or emailed as an online link after your call), so we can't fill it in for you. Copy or paste each answer into the matching box. If you are filling in a paper form, you can print the extra sheets we've made (one per question), write your NI number on each, and send them with it. Write \'see extra sheet\' in any box that is too small.</p>"
        send=""
        if is_full(tier):
            send=f"<h2>Phone call and assessment prep</h2><ul>{li(draft.get('call_prep',[]))}</ul><p>PIP new claims line: 0800 917 2222 (Mon–Fri 8am–5pm). Return the form within 1 month, or call the PIP enquiry line if you need more time.</p>"
        title="Your PIP answers" if form=="pip" else "Your PIP review (AR1) answers"
    return f"""<!doctype html><html lang="en-GB"><meta charset="utf-8"><title>{title}</title>{CSS}
<h1>{title}</h1><p>Prepared by {e(brand)}.</p>
<div class="box"><strong>Please check every answer.</strong> It must be true. Change anything that isn't quite right. You send the form yourself. We can't guarantee an award; DWP decides. Giving false information is an offence.</div>
{delivery}<h2>Things to check</h2><ul>{li(draft.get('notes_for_user',[]))}</ul>
<h2>Your answers</h2>{rows}{extra}
<h2>Evidence checklist</h2><ul>{li(draft.get('checklist',[]))}</ul>
<h2>Draft letter to the GP or consultant</h2><pre style="white-space:pre-wrap;font:inherit">{e(draft.get('gp_letter',''))}</pre>
<h2>Draft statement from someone who knows {'the child' if form=='dla' else 'you'}</h2><p>{e(draft.get('carer_statement',''))}</p>{send}
<p style="font-size:.85em;color:#555">Independent service, not affiliated with DWP.</p></html>"""

# ---- Access to Work ----
from facts import v as _v
def atw_sections(draft, tier):
    F=FORMS["atw"]; a=draft.get("answers",{}); out=[]
    for i,t,_ in F["topics"]:
        out.append(("q",t,a.get(i,"") or "Not discussed. Add this in your own words if it applies."))
    if is_full(tier):
        out.append(("h","Support-needs summary",[f"{x.get('item','')} ({x.get('type','')}) – cost: {x.get('cost','[quote]')}" for x in draft.get("support_summary",[])]))
        out.append(("table","Cost list (fill in quotes as you get them)",draft.get("support_summary",[])))
        out.append(("h","Before you call or apply",draft.get("call_prep",[])))
        import deadlines
        out.append(("h","If you disagree with the decision: asking for a reconsideration",deadlines.atw_reconsider_script()))
    out.append(("h","Evidence and details to have ready",draft.get("checklist",[])))
    return out
ATW_TOP=("Answers in the order the Access to Work application asks about them. Paste or read these out when you apply online or by phone. "
 "Please check every answer is true and in your words; change anything that isn't right. You add your personal details, workplace address, workplace contact and NI number yourself. "
 "Access to Work decides what support to offer; we can't guarantee an outcome.")
def pack_atw(brand, draft, tier):
    e=_esc; body=""
    for kind,t,x in atw_sections(draft,tier):
        if kind=="q": body+=f"<h3>{e(t)}</h3><p>{e(x)}</p>"
        elif kind=="table": body+=f"<h2>{e(t)}</h2><table border=1 cellpadding=6 style='border-collapse:collapse'><tr><th>Item</th><th>Type</th><th>Cost</th></tr>"+"".join(f"<tr><td>{e(r.get('item',''))}</td><td>{e(r.get('type',''))}</td><td>{e(r.get('cost','[quote]'))}</td></tr>" for r in x)+"</table>"
        else: body+=f"<h2>{e(t)}</h2><ul>"+"".join(f"<li>{e(i)}</li>" for i in x)+"</ul>"
    return f"""<!doctype html><html lang="en-GB"><meta charset="utf-8"><title>Your Access to Work answers</title>{CSS}
<h1>Your Access to Work answers</h1><p>Prepared by {e(brand)}.</p><div class=box>{e(ATW_TOP)}</div><h2>Your answers</h2>{body}
<p style="font-size:.85em;color:#555">Independent service, not affiliated with DWP or Access to Work.</p></html>"""
def atw_docx(brand, draft, tier, path):
    from docx import Document
    d=Document(); d.add_heading("Your Access to Work answers",0); d.add_paragraph(f"Prepared by {brand}."); d.add_paragraph(ATW_TOP)
    d.add_heading("Your answers",1)
    for kind,t,x in atw_sections(draft,tier):
        if kind=="q": d.add_heading(t,2); d.add_paragraph(x)
        elif kind=="table":
            d.add_heading(t,1); tb=d.add_table(rows=1,cols=3); tb.style="Table Grid"
            for c,h in zip(tb.rows[0].cells,("Item","Type","Cost")): c.text=h
            for r in x:
                cs=tb.add_row().cells; cs[0].text=r.get("item",""); cs[1].text=r.get("type",""); cs[2].text=r.get("cost","[quote]")
        else:
            d.add_heading(t,1)
            for i in x: d.add_paragraph(i,style="List Bullet")
    d.add_paragraph(f"Independent service, not affiliated with DWP or Access to Work.")
    d.save(path)
